"""Text extraction from CV / job-description uploads (PDF, DOCX, TXT).

Extraction is deliberately conservative: it preserves the original wording and
the Azerbaijani characters, and never rewrites or summarises. Anything that
reaches the model is text the candidate actually wrote.
"""
from __future__ import annotations

import logging
import unicodedata
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = (".pdf", ".docx", ".txt", ".md")

#: Guards against a pathological upload exhausting memory or the token budget.
MAX_CHARACTERS = 400_000


class DocumentError(RuntimeError):
    """Raised when a document cannot be read or is of an unsupported type."""


@dataclass(frozen=True)
class ExtractedDocument:
    path: str
    kind: str          # "cv" | "job_description"
    text: str
    page_count: int
    char_count: int

    @property
    def name(self) -> str:
        return Path(self.path).name


def extract_text(path: str | Path, kind: str = "cv") -> ExtractedDocument:
    """Read a document and return its plain text."""
    source = Path(path)
    if not source.exists():
        raise DocumentError(f"File not found: {source}")
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentError(
            f"Unsupported file type {suffix!r}. Use one of: PDF, DOCX, TXT."
        )

    if suffix == ".pdf":
        text, pages = _extract_pdf(source)
    elif suffix == ".docx":
        text, pages = _extract_docx(source)
    else:
        text, pages = _extract_txt(source), 1

    text = _clean(text)
    if not text.strip():
        raise DocumentError(
            f"No text could be extracted from {source.name}. "
            "If it is a scanned PDF it must be run through OCR first."
        )
    if len(text) > MAX_CHARACTERS:
        log.warning("Truncating %s from %d to %d characters", source.name, len(text), MAX_CHARACTERS)
        text = text[:MAX_CHARACTERS]

    return ExtractedDocument(
        path=str(source), kind=kind, text=text, page_count=pages, char_count=len(text)
    )


def _extract_pdf(source: Path) -> tuple[str, int]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("pypdf is required to read PDF files") from exc
    try:
        reader = PdfReader(str(source))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise DocumentError(f"{source.name} is password protected") from exc
        pages = [(page.extract_text() or "") for page in reader.pages]
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError(f"Could not read {source.name}: {exc}") from exc
    return "\n\n".join(pages), len(pages)


def _extract_docx(source: Path) -> tuple[str, int]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("python-docx is required to read DOCX files") from exc
    try:
        document = docx.Document(str(source))
    except Exception as exc:
        raise DocumentError(f"Could not read {source.name}: {exc}") from exc

    parts: list[str] = [p.text for p in document.paragraphs]
    # Skills and experience are very often laid out in tables.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), 1


def _extract_txt(source: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1254", "cp1251", "latin-1"):
        try:
            return source.read_text(encoding=encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    raise DocumentError(f"Could not decode {source.name} as text")


def _clean(text: str) -> str:
    """Normalise whitespace while preserving Azerbaijani characters verbatim."""
    text = unicodedata.normalize("NFC", text or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    lines = [line.rstrip() for line in text.split("\n")]

    out: list[str] = []
    blanks = 0
    for line in lines:
        if line.strip():
            blanks = 0
            out.append(" ".join(line.split()))
        else:
            blanks += 1
            if blanks <= 1:          # collapse runs of blank lines
                out.append("")
    return "\n".join(out).strip()
